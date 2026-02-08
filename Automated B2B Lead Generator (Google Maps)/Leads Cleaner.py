import pandas as pd
import re
import os
import glob
import numpy as np

# --- CONFIGURATION ---
LEADS_FOLDER_NAME = 'Leads'
OUTPUT_FOLDER_NAME = 'Cleaned'

# --- 1. SETUP FOLDERS ---
base_dir = os.getcwd()
leads_path = os.path.join(base_dir, LEADS_FOLDER_NAME)
output_path = os.path.join(base_dir, OUTPUT_FOLDER_NAME)

if not os.path.exists(output_path):
    os.makedirs(output_path)

# --- 2. FIND FILES ---
def get_files():
    search_path = leads_path if os.path.exists(leads_path) else base_dir
    files = glob.glob(os.path.join(search_path, "*.csv")) + glob.glob(os.path.join(search_path, "*.xlsx"))
    return files

# --- 3. CLEANING FUNCTIONS ---
def clean_email_cell(cell_value):
    if pd.isna(cell_value):
        return cell_value

    # Extract emails from messy strings like:
    # "Email: info@site.com (preferred)", "info@site.com | sales@site.com", etc.
    text = str(cell_value).strip()
    if not text:
        return ""

    email_pattern = re.compile(r"(?i)[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}")
    found_emails = email_pattern.findall(text)

    if not found_emails:
        return text

    # "Ghost Words" to remove from end of email
    bad_suffixes = [
        'monday', 'opening', 'full', 'head', 'follow',
        'hoursmon', 'need', 'our', 'contact'
    ]

    cleaned_emails = []
    for email in found_emails:
        email = email.strip().strip("<>()[]{}\"'.,;:! ")

        if email.lower().startswith("mailto:"):
            email = email[7:].strip()

        # Remove "email" prefix only if it becomes a valid email
        if email.lower().startswith('email') and not email.lower().startswith('email@'):
            candidate = re.sub(r'^email', '', email, flags=re.IGNORECASE).lstrip(':').strip()
            if email_pattern.fullmatch(candidate):
                email = candidate

        # Remove leading digits only if it becomes a valid email
        candidate = re.sub(r'^\d+', '', email)
        if candidate != email and email_pattern.fullmatch(candidate):
            email = candidate

        # Remove known bad suffixes only if it becomes a valid email
        for suffix in bad_suffixes:
            if email.lower().endswith(suffix):
                candidate = re.sub(f'{suffix}$', '', email, flags=re.IGNORECASE)
                if email_pattern.fullmatch(candidate):
                    email = candidate

        cleaned_emails.append(email.lower())

    # Deduplicate while preserving order
    unique_emails = []
    seen = set()
    for email in cleaned_emails:
        if email and email not in seen:
            seen.add(email)
            unique_emails.append(email)

    return ', '.join(unique_emails)

def auto_adjust_excel_columns(file_path, sheet_name=None, min_width=10, max_width=60, padding=2):
    """
    Expands Excel column widths to fit text size (approximation).

    Note: openpyxl can't do true Excel "AutoFit", so this estimates width based
    on the longest visible line in each column and caps it to avoid huge columns.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    for col_idx in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col_idx).value
        header = str(header_val).strip() if header_val is not None else ""
        header_lower = header.lower()

        # Per-column caps (keeps links / long notes reasonable)
        col_max_width = max_width
        if "sales angle" in header_lower or "audit issues" in header_lower:
            col_max_width = max(max_width, 80)
        elif "link" in header_lower or "website" in header_lower:
            col_max_width = min(col_max_width, 45)

        max_len = len(header)
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue

            text = str(value)
            longest_line = max((len(line) for line in text.splitlines()), default=len(text))
            if longest_line > max_len:
                max_len = longest_line

        adjusted = max(min_width, min(col_max_width, max_len + padding))
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted

    wb.save(file_path)

def fill_missing_audits(df):
    # Ensure columns exist
    if 'Audit Issues' not in df.columns:
        df['Audit Issues'] = np.nan
    if 'Sales Angle' not in df.columns:
        df['Sales Angle'] = np.nan
    if 'Website' not in df.columns:
        df['Website'] = np.nan

    # Case A: Has Website but Missing Audit
    # We fill it with a generic "Manual Check" label so the cell isn't blank
    mask_has_site = (df['Audit Issues'].isna() | (df['Audit Issues'] == '')) & (df['Website'].notna()) & (df['Website'] != '')
    
    df.loc[mask_has_site, 'Audit Issues'] = "Manual Audit Required (Site Active)"
    df.loc[mask_has_site, 'Sales Angle'] = "Website is live but needs a professional review for conversion gaps."

    # Case B: No Website and Missing Audit
    mask_no_site = (df['Audit Issues'].isna() | (df['Audit Issues'] == '')) & (df['Website'].isna() | (df['Website'] == ''))
    
    df.loc[mask_no_site, 'Audit Issues'] = "No Website Found"
    df.loc[mask_no_site, 'Sales Angle'] = "You are invisible to digital clients."
    
    return df

# --- 4. MAIN INTERFACE ---
files = get_files()

if not files:
    print("❌ No CSV or Excel files found in the 'Leads' folder!")
else:
    print(f"--- Found {len(files)} files ---")
    for i, f in enumerate(files):
        print(f"[{i+1}] {os.path.basename(f)}")
    
    selected_file = None
    while True:
        try:
            choice = input("\nEnter the number of the file to clean (or 'q' to quit): ")
            if choice.lower() == 'q':
                break
            
            idx = int(choice) - 1
            if 0 <= idx < len(files):
                selected_file = files[idx]
                break
            else:
                print("Invalid number. Try again.")
        except ValueError:
            print("Please enter a number.")
            
    if selected_file:
        print(f"\nProcessing: {os.path.basename(selected_file)}...")
        try:
            # Load
            if selected_file.endswith('.csv'):
                try:
                    df = pd.read_csv(selected_file)
                except:
                    df = pd.read_csv(selected_file, encoding='latin1')
            else:
                df = pd.read_excel(selected_file)

            # 1. Clean Emails
            if 'Business Email' in df.columns:
                df['Business Email'] = df['Business Email'].apply(clean_email_cell)
                print("✅ Email column cleaned.")
            
            # 2. Fill Missing Audits
            df = fill_missing_audits(df)
            print("✅ Missing Audit Data filled.")

            # Save
            filename = os.path.basename(selected_file)
            base_name, _ = os.path.splitext(filename)
            save_name = f"Cleaned_{base_name}.xlsx"
                
            save_path = os.path.join(output_path, save_name)
            df.to_excel(save_path, index=False)

            # Auto-adjust column widths so text is readable in Excel
            try:
                auto_adjust_excel_columns(save_path)
                print("✅ Excel column widths adjusted.")
            except Exception as e:
                print(f"⚠️ Warning: couldn't adjust Excel column widths: {e}")

            print(f"🎉 Success! File saved to: {OUTPUT_FOLDER_NAME}/{save_name}")
            
        except Exception as e:
            print(f"❌ Error: {e}")
